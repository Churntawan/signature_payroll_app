import openpyxl
import datetime
import urllib.request
import json

EXCEL_PATH = r"C:\Users\User\Desktop\Signature_Payroll_V3_Standard.xlsx"
SUPABASE_URL = "https://qsmigegcefcbohmufywh.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InFzbWlnZWdjZWZjYm9obXVmeXdoIiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODkwNDM4MDAsImV4cCI6MjEwNDYxOTgwMH0.jIEkmGeSVzht81fGEQnxOM-n9TGG7AFumkFEe5SVGTk"

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

def post_batch(table, items):
    if not items:
        return
    url = f"{SUPABASE_URL}/rest/v1/{table}"
    data = json.dumps(items).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=data,
        headers={
            "apikey": SUPABASE_KEY,
            "Authorization": f"Bearer {SUPABASE_KEY}",
            "Content-Type": "application/json",
            "Prefer": "return=minimal",
        },
    )
    with urllib.request.urlopen(req) as res:
        pass

def format_date(val):
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%Y-%m-%d")
    if val:
        return str(val).strip()[:10]
    return None

# 1. Fetch valid employee codes from Supabase
emp_req = urllib.request.Request(
    f"{SUPABASE_URL}/rest/v1/employees?select=ep_code",
    headers={"apikey": SUPABASE_KEY, "Authorization": f"Bearer {SUPABASE_KEY}"}
)
with urllib.request.urlopen(emp_req) as res:
    valid_eps = {x["ep_code"] for x in json.loads(res.read())}

print(f"Valid employees in Supabase: {len(valid_eps)}")

# 2. Migrate Attendance_Log
att_sheet = wb["Attendance_Log"]
att_items = []
count_att = 0
for r in list(att_sheet.iter_rows(min_row=2, values_only=True)):
    if r and r[0] and r[2]:
        ep = str(r[2]).strip()
        if ep not in valid_eps:
            continue
        d = format_date(r[1])
        if not d:
            continue
        att_items.append({
            "log_id": str(r[0]).strip(),
            "date": d,
            "ep_code": ep,
            "nickname": str(r[3]).strip() if r[3] else "",
            "category": str(r[4]).strip() if r[4] else "Day-off",
            "shift": str(r[5]).strip() if r[5] else "Normal",
            "units": float(r[6]) if r[6] is not None else 1.0,
            "note": str(r[7]).strip() if len(r) > 7 and r[7] else "",
        })
        if len(att_items) >= 200:
            post_batch("attendance_log", att_items)
            count_att += len(att_items)
            print(f"Uploaded {count_att} attendance records...")
            att_items = []

if att_items:
    post_batch("attendance_log", att_items)
    count_att += len(att_items)
    print(f"Uploaded total {count_att} attendance records!")

# 3. Migrate Payroll_Adjustments
adj_sheet = wb["Payroll_Adjustments"]
adj_items = []
count_adj = 0
for r in list(adj_sheet.iter_rows(min_row=2, values_only=True)):
    if r and r[0] and r[3]:
        ep = str(r[3]).strip()
        if ep not in valid_eps:
            continue
        adj_items.append({
            "adj_id": str(r[0]).strip(),
            "period": str(r[1]).strip() if r[1] else "",
            "due_date": format_date(r[2]),
            "ep_code": ep,
            "nickname": str(r[4]).strip() if r[4] else "",
            "type": str(r[5]).strip() if r[5] else "Deduction",
            "category": str(r[6]).strip() if r[6] else "Advance",
            "description": str(r[7]).strip() if r[7] else "",
            "amount": float(r[8]) if r[8] is not None else 0.0,
            "status": str(r[9]).strip() if len(r) > 9 and r[9] else "Pending",
        })
        if len(adj_items) >= 200:
            post_batch("payroll_adjustments", adj_items)
            count_adj += len(adj_items)
            print(f"Uploaded {count_adj} adjustment records...")
            adj_items = []

if adj_items:
    post_batch("payroll_adjustments", adj_items)
    count_adj += len(adj_items)
    print(f"Uploaded total {count_adj} adjustment records!")

print("All migration completed successfully!")
