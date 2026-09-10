import os
import openpyxl
from datetime import datetime, date
from typing import Optional, List
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

EXCEL_PATH = r"C:\Users\User\Desktop\Signature_Payroll_V3_Standard.xlsx"

app = FastAPI(title="Signature Payroll API", version="2.0.0")

# Enable CORS for Flutter web / desktop
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def get_workbook():
    if not os.path.exists(EXCEL_PATH):
        raise HTTPException(status_code=500, detail=f"Database file not found: {EXCEL_PATH}")
    return openpyxl.load_workbook(EXCEL_PATH)

def format_date(val):
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    return str(val) if val is not None else ""

# -------------------------------------------------------------
# PYDANTIC SCHEMAS
# -------------------------------------------------------------
class AttendanceItem(BaseModel):
    date: str # YYYY-MM-DD
    ep_code: str
    nickname: str
    category: str # 'Day-off', 'Sick', 'Half-day', 'OT Days', 'Work Days'
    shift: Optional[str] = "Normal"
    units: Optional[float] = 1.0
    note: Optional[str] = ""

class AdjustmentItem(BaseModel):
    period: str # e.g. '2025-01'
    due_date: str # YYYY-MM-DD
    ep_code: str
    nickname: str
    type: str # 'Income' or 'Deduction'
    category: str # 'Advance', 'Work Permit', 'Passport / CI', 'Bonus', 'Other'
    description: Optional[str] = ""
    amount: float
    status: Optional[str] = "Pending"

class EmployeeItem(BaseModel):
    ep_code: str
    nickname: str
    status: str # 'Active' or 'Resigned'
    base_salary: float
    pay_group: str # 'Date : 1', 'Date : 10', 'Date : 20'
    stay_outside: Optional[str] = "No"
    start_date: Optional[str] = None
    resign_date: Optional[str] = None
    note: Optional[str] = ""

class PayrollSummaryItem(BaseModel):
    payroll_id: Optional[str] = None
    period: str # e.g. '2025-01'
    ep_code: str
    nickname: str
    pay_type: Optional[str] = "Full Month"
    base_salary: float
    work_days: int = 26
    day_off: int = 4
    sick: int = 0
    half_day: int = 0
    ot_days: int = 0
    base_pay: float
    total_extra: float = 0.0
    total_deduction: float = 0.0
    net_pay: float
    status: Optional[str] = "Approved"
    note: Optional[str] = ""

class PayrollSummaryBatch(BaseModel):
    records: List[PayrollSummaryItem]


# -------------------------------------------------------------
# API ENDPOINTS
# -------------------------------------------------------------

@app.get("/api/health")
def health():
    return {
        "status": "online",
        "database": EXCEL_PATH,
        "exists": os.path.exists(EXCEL_PATH)
    }

@app.get("/api/periods")
def get_periods():
    wb = get_workbook()
    sheet = wb["Config_Lists"]
    periods = []
    # Column 9 is Periods
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row and len(row) >= 9 and row[8]:
            p = str(row[8]).strip()
            if p and p not in periods:
                periods.append(p)
    return {"periods": periods}

@app.get("/api/employees")
def get_employees():
    wb = get_workbook()
    sheet = wb["Employees"]
    employees = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            employees.append({
                "ep_code": str(row[0]).strip(),
                "nickname": str(row[1]).strip() if row[1] else "",
                "status": str(row[2]).strip() if row[2] else "Active",
                "base_salary": float(row[3]) if row[3] is not None else 0.0,
                "pay_group": str(row[4]).strip() if row[4] else "Date : 10",
                "stay_outside": str(row[5]).strip() if row[5] else "No",
                "start_date": format_date(row[6]),
                "resign_date": format_date(row[7]),
                "note": str(row[8]).strip() if len(row) > 8 and row[8] else "",
            })
    return {"employees": employees}

@app.get("/api/attendance")
def get_attendance(
    period: Optional[str] = None,
    ep_code: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
):
    wb = get_workbook()
    sheet = wb["Attendance_Log"]
    logs = []

    # Calculate cycle window if period is given (e.g. 2025-01 spans 2024-12-02 to 2025-01-20)
    cycle_start = start_date
    cycle_end = end_date
    if period and not (cycle_start and cycle_end) and len(period) >= 7:
        try:
            parts = period.split("-")
            y = int(parts[0])
            m = int(parts[1])
            prev_y = y if m > 1 else y - 1
            prev_m = m - 1 if m > 1 else 12
            cycle_start = f"{prev_y:04d}-{prev_m:02d}-02"
            cycle_end = f"{y:04d}-{m:02d}-20"
        except Exception:
            pass

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            log_date = format_date(row[1])
            curr_ep = str(row[2]).strip() if row[2] else ""
            if ep_code and curr_ep != ep_code:
                continue
            if cycle_start and cycle_end:
                if not (cycle_start <= log_date <= cycle_end):
                    continue
            elif period and not log_date.startswith(period[:7]):
                continue

            logs.append({
                "log_id": str(row[0]).strip(),
                "date": log_date,
                "ep_code": curr_ep,
                "nickname": str(row[3]).strip() if row[3] else "",
                "category": str(row[4]).strip() if row[4] else "Day-off",
                "shift": str(row[5]).strip() if row[5] else "Normal",
                "units": float(row[6]) if row[6] is not None else 1.0,
                "note": str(row[7]).strip() if len(row) > 7 and row[7] else "",
            })
    return {"attendance": logs}

@app.post("/api/attendance")
def create_attendance(item: AttendanceItem):
    wb = get_workbook()
    sheet = wb["Attendance_Log"]
    # Generate next Log ID
    next_id = sheet.max_row
    log_id = f"LOG{next_id:04d}"
    
    sheet.append([
        log_id,
        item.date,
        item.ep_code,
        item.nickname,
        item.category,
        item.shift,
        item.units,
        item.note,
    ])
    wb.save(EXCEL_PATH)
    return {"message": "Attendance record created", "log_id": log_id, "data": item.dict()}

@app.get("/api/adjustments")
def get_adjustments(period: Optional[str] = None, ep_code: Optional[str] = None):
    wb = get_workbook()
    sheet = wb["Payroll_Adjustments"]
    adjustments = []
    # Columns: Adj ID, Period, Due Date, EP Code, Nickname, Type, Category, Description, Amount, Status
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            curr_period = str(row[1]).strip() if row[1] else ""
            curr_ep = str(row[3]).strip() if row[3] else ""
            if period and curr_period != period:
                continue
            if ep_code and curr_ep != ep_code:
                continue
            adjustments.append({
                "adj_id": str(row[0]).strip(),
                "period": curr_period,
                "due_date": format_date(row[2]),
                "ep_code": curr_ep,
                "nickname": str(row[4]).strip() if row[4] else "",
                "type": str(row[5]).strip() if row[5] else "Deduction",
                "category": str(row[6]).strip() if row[6] else "Advance",
                "description": str(row[7]).strip() if row[7] else "",
                "amount": float(row[8]) if row[8] is not None else 0.0,
                "status": str(row[9]).strip() if len(row) > 9 and row[9] else "Pending",
            })
    return {"adjustments": adjustments}

@app.post("/api/adjustments")
def create_adjustment(item: AdjustmentItem):
    wb = get_workbook()
    sheet = wb["Payroll_Adjustments"]
    next_id = sheet.max_row
    adj_id = f"ADJ{next_id:04d}"
    
    sheet.append([
        adj_id,
        item.period,
        item.due_date,
        item.ep_code,
        item.nickname,
        item.type,
        item.category,
        item.description,
        item.amount,
        item.status,
    ])
    wb.save(EXCEL_PATH)
    return {"message": "Adjustment record created", "adj_id": adj_id, "data": item.dict()}

@app.get("/api/payroll-summary")
def get_payroll_summary(period: Optional[str] = None, ep_code: Optional[str] = None):
    wb = get_workbook()
    sheet = wb["Payroll_Summary"]
    summaries = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            curr_period = str(row[1]).strip() if row[1] else ""
            curr_ep = str(row[2]).strip() if row[2] else ""
            if period and curr_period != period:
                continue
            if ep_code and curr_ep != ep_code:
                continue
            summaries.append({
                "payroll_id": str(row[0]).strip(),
                "period": curr_period,
                "ep_code": curr_ep,
                "nickname": str(row[3]).strip() if row[3] else "",
                "pay_type": str(row[4]).strip() if row[4] else "Full Month",
                "base_salary": float(row[5]) if row[5] is not None else 0.0,
                "work_days": int(row[6]) if row[6] is not None else 26,
                "day_off": int(row[7]) if row[7] is not None else 4,
                "sick": int(row[8]) if row[8] is not None else 0,
                "half_day": int(row[9]) if row[9] is not None else 0,
                "ot_days": int(row[10]) if row[10] is not None else 0,
                "base_pay": float(row[11]) if row[11] is not None else 0.0,
                "total_extra": float(row[12]) if row[12] is not None else 0.0,
                "total_deduction": float(row[13]) if row[13] is not None else 0.0,
                "net_pay": float(row[14]) if row[14] is not None else 0.0,
                "status": str(row[15]).strip() if len(row) > 15 and row[15] else "Approved",
                "note": str(row[16]).strip() if len(row) > 16 and row[16] else "",
            })
    return {"summaries": summaries}

@app.post("/api/payroll-summary")
def save_payroll_summary(batch: PayrollSummaryBatch):
    wb = get_workbook()
    sheet = wb["Payroll_Summary"]

    existing_map = {}
    for idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), start=2):
        if row and len(row) >= 3 and row[1].value and row[2].value:
            p = str(row[1].value).strip()
            ep = str(row[2].value).strip()
            existing_map[(p, ep)] = idx

    updated_count = 0
    created_count = 0

    for item in batch.records:
        key = (item.period.strip(), item.ep_code.strip())
        if key in existing_map:
            row_idx = existing_map[key]
            sheet.cell(row=row_idx, column=4, value=item.nickname)
            sheet.cell(row=row_idx, column=5, value=item.pay_type)
            sheet.cell(row=row_idx, column=6, value=item.base_salary)
            sheet.cell(row=row_idx, column=7, value=item.work_days)
            sheet.cell(row=row_idx, column=8, value=item.day_off)
            sheet.cell(row=row_idx, column=9, value=item.sick)
            sheet.cell(row=row_idx, column=10, value=item.half_day)
            sheet.cell(row=row_idx, column=11, value=item.ot_days)
            sheet.cell(row=row_idx, column=12, value=item.base_pay)
            sheet.cell(row=row_idx, column=13, value=item.total_extra)
            sheet.cell(row=row_idx, column=14, value=item.total_deduction)
            sheet.cell(row=row_idx, column=15, value=item.net_pay)
            sheet.cell(row=row_idx, column=16, value=item.status)
            sheet.cell(row=row_idx, column=17, value=item.note)
            updated_count += 1
        else:
            next_num = sheet.max_row
            payroll_id = f"STO{next_num:03d}"
            sheet.append([
                payroll_id,
                item.period,
                item.ep_code,
                item.nickname,
                item.pay_type,
                item.base_salary,
                item.work_days,
                item.day_off,
                item.sick,
                item.half_day,
                item.ot_days,
                item.base_pay,
                item.total_extra,
                item.total_deduction,
                item.net_pay,
                item.status,
                item.note,
            ])
            created_count += 1
            existing_map[key] = sheet.max_row

    wb.save(EXCEL_PATH)
    return {
        "message": "Payroll summary saved to Excel successfully",
        "updated": updated_count,
        "created": created_count,
        "total": len(batch.records),
    }

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)

